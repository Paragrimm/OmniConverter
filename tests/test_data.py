import json
import tomllib

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from omniconverter.backends.data import coerce, flatten, records_to_table
from omniconverter.core.errors import ConversionError
from omniconverter.core.formats import get_format


def convert(converter, path, target, **options):
    source = converter.inspect(path)
    fmt = get_format(target)
    out = path.with_name(f"{path.stem}-out.{fmt.extension}")
    return converter.convert(source, fmt, out, options)


def test_semicolon_csv_to_json_detects_numbers(converter, tmp_path):
    src = tmp_path / "people.csv"
    src.write_text("name;age;zip\nAnna;31;01234\nBen;4.5;99999\n", encoding="utf-8")
    data = json.loads(convert(converter, src, "json").read_text(encoding="utf-8"))
    assert data == [{"name": "Anna", "age": 31, "zip": "01234"},
                    {"name": "Ben", "age": 4.5, "zip": 99999}]


def test_csv_numbers_can_stay_text(converter, tmp_path):
    src = tmp_path / "p.csv"
    src.write_text("a,b\n1,2\n", encoding="utf-8")
    data = json.loads(convert(converter, src, "json", detect_numbers=False).read_text())
    assert data == [{"a": "1", "b": "2"}]


def test_windows_1252_csv(converter, tmp_path):
    src = tmp_path / "legacy.csv"
    src.write_bytes("stadt,einwohner\nMünchen,1500000\n".encode("cp1252"))
    data = yaml.safe_load(convert(converter, src, "yaml").read_text(encoding="utf-8"))
    assert data == [{"stadt": "München", "einwohner": 1500000}]


def test_nested_json_to_csv_is_flattened(converter, tmp_path):
    src = tmp_path / "d.json"
    src.write_text(json.dumps([{"id": 1, "user": {"name": "x", "tags": ["a"]}},
                               {"id": 2, "active": True}]))
    out = convert(converter, src, "csv", delimiter="semicolon", excel_bom=True)
    raw = out.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    lines = raw.decode("utf-8-sig").splitlines()
    assert lines[0] == "id;user.name;user.tags;active"
    assert lines[1] == '1;x;"[""a""]";'
    assert lines[2] == "2;;;true"


def test_xlsx_roundtrip(converter, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["name", "amount"])
    ws.append(["a", 1.5])
    ws.append(["b", 2])
    extra = wb.create_sheet("Other")
    extra.append(["x"])
    extra.append([42])
    src = tmp_path / "book.xlsx"
    wb.save(src)

    rows = convert(converter, src, "csv").read_text().splitlines()
    assert rows == ["name,amount", "a,1.5", "b,2"]
    other = json.loads(convert(converter, src, "json", sheet="Other").read_text())
    assert other == [{"x": 42}]
    with pytest.raises(ConversionError):
        convert(converter, src, "json", sheet="Missing")

    back = convert(converter, tmp_path / "book-out.csv", "xlsx")
    ws2 = load_workbook(back).active
    assert [c.value for c in ws2[1]] == ["name", "amount"]


def test_structured_conversions(converter, tmp_path):
    src = tmp_path / "config.yaml"
    src.write_text("server:\n  host: localhost\n  port: 8080\nfeatures: [a, b]\n")
    toml_out = convert(converter, src, "toml")
    assert tomllib.loads(toml_out.read_text()) == {
        "server": {"host": "localhost", "port": 8080}, "features": ["a", "b"]}
    json_out = convert(converter, toml_out, "json", indent=0)
    assert "\n" not in json_out.read_text().strip()


def test_table_to_toml_and_back(converter, tmp_path):
    src = tmp_path / "t.csv"
    src.write_text("k,v\na,1\nb,\n")
    toml_out = convert(converter, src, "toml")
    assert tomllib.loads(toml_out.read_text()) == {"rows": [{"k": "a", "v": 1}, {"k": "b"}]}
    back = convert(converter, toml_out, "csv")
    assert back.read_text().splitlines() == ["k,v", "a,1", "b,"]


def test_non_tabular_data_is_rejected(converter, tmp_path):
    src = tmp_path / "list.json"
    src.write_text("[1, 2, 3]")
    with pytest.raises(ConversionError) as err:
        convert(converter, src, "csv")
    assert "tabular" in err.value.message


def test_invalid_json_message(converter, tmp_path):
    src = tmp_path / "bad.json"
    src.write_text("{nope")
    with pytest.raises(ConversionError):
        convert(converter, src, "yaml")


@pytest.mark.parametrize(("raw", "value"), [
    ("42", 42), ("-3.5", -3.5), ("1e3", 1000.0), ("007", "007"), ("+49", "+49"), ("", None),
    ("abc", "abc"),
])
def test_coerce(raw, value):
    assert coerce(raw) == value


def test_flatten_and_records():
    assert flatten({"a": {"b": {"c": 1}}, "d": [1]}) == {"a.b.c": 1, "d": "[1]"}
    header, rows = records_to_table({"only": [{"x": 1}, {"y": 2}]})
    assert header == ["x", "y"] and rows == [[1, None], [None, 2]]

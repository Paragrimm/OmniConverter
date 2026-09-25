"""Tabular and structured data: CSV, TSV, XLSX, JSON, YAML, TOML (pure Python)."""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import tomllib
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from omniconverter.core.backend import Backend, Conversion, ConversionContext, ConversionRequest
from omniconverter.core.errors import ConversionError
from omniconverter.core.formats import Format
from omniconverter.core.options import Kind, Option
from omniconverter.core.probe import MediaInfo
from omniconverter.i18n import t

TABULAR = ("csv", "tsv", "xlsx")
STRUCTURED = ("json", "yaml", "toml")
FORMATS = TABULAR + STRUCTURED

_DELIMITERS = {"comma": ",", "semicolon": ";", "tab": "\t"}
_NUMBER = re.compile(r"-?(0|[1-9]\d*)(\.\d+)?([eE][+-]?\d+)?")


class DataBackend(Backend):
    id = "data"

    def conversions(self) -> Iterable[Conversion]:
        for src in FORMATS:
            for dst in FORMATS:
                if src != dst:
                    # Beats LibreOffice for csv ↔ xlsx: faster and needs nothing installed.
                    yield Conversion(src, dst, priority=80)

    def options(self, source: Format, target: Format, media: MediaInfo | None = None
                ) -> list[Option]:
        opts: list[Option] = []
        if source.id == "xlsx":
            opts.append(Option("sheet", t("opt.sheet"), Kind.TEXT, "", help=t("opt.sheet_help")))
        if source.id in TABULAR and target.id in STRUCTURED:
            opts.append(Option("detect_numbers", t("opt.detect_numbers"), Kind.BOOL, True))
        if target.id == "csv":
            opts.append(Option(
                "delimiter", t("opt.delimiter"), Kind.CHOICE, "comma",
                choices=tuple((k, t(f"opt.delimiter.{k}")) for k in _DELIMITERS),
            ))
        if target.id in ("csv", "tsv"):
            opts.append(Option("excel_bom", t("opt.excel_bom"), Kind.BOOL, False))
        if target.id == "json":
            opts.append(Option("indent", t("opt.indent"), Kind.CHOICE, 2,
                               choices=((2, "2"), (4, "4"), (0, t("opt.indent.compact")))))
        return opts

    def convert(self, request: ConversionRequest, output: Path, ctx: ConversionContext) -> None:
        ctx.progress(None)
        src, dst, opts = request.source_format.id, request.target_format.id, request.options
        try:
            if src in TABULAR:
                header, rows = read_table(request.source, src, opts.get("sheet", ""))
                if dst in TABULAR:
                    write_table(output, dst, header, rows, opts)
                    return
                data: Any = table_to_records(header, rows, opts.get("detect_numbers", True))
            else:
                data = read_structured(request.source, src)
                if dst in TABULAR:
                    header, rows = records_to_table(data)
                    write_table(output, dst, header, rows, opts)
                    return
            write_structured(output, dst, data, opts)
        except ConversionError:
            raise
        except (OSError, ValueError, TypeError, KeyError) as exc:
            raise ConversionError(t("error.data_read", problem=str(exc)), repr(exc)) from exc


# -- reading -------------------------------------------------------------------------------


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def read_table(path: Path, fmt: str, sheet: str = "") -> tuple[list[str], list[list[Any]]]:
    if fmt == "xlsx":
        return _read_xlsx(path, sheet)
    text = _read_text(path)
    if fmt == "tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter) if row]
    if not rows:
        return [], []
    return rows[0], rows[1:]


def _read_xlsx(path: Path, sheet: str) -> tuple[list[str], list[list[Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet:
            if sheet not in wb.sheetnames:
                raise ConversionError(t("error.sheet_missing", sheet=sheet))
            ws = wb[sheet]
        else:
            ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    # Drop trailing empty rows/columns that spreadsheets like to carry around.
    rows = [r for r in rows if any(v not in (None, "") for v in r)]
    if not rows:
        return [], []
    width = max(max((i + 1 for i, v in enumerate(r) if v not in (None, "")), default=0)
                for r in rows)
    rows = [(r + [None] * width)[:width] for r in rows]
    header = [str(v) if v is not None else f"column{i + 1}" for i, v in enumerate(rows[0])]
    return header, rows[1:]


def read_structured(path: Path, fmt: str) -> Any:
    text = _read_text(path)
    if fmt == "json":
        return json.loads(text)
    if fmt == "yaml":
        import yaml

        docs = list(yaml.safe_load_all(text))
        return docs[0] if len(docs) == 1 else docs
    return tomllib.loads(text)


# -- shaping -------------------------------------------------------------------------------


def coerce(value: Any) -> Any:
    """Turn numeric strings into numbers (keeping e.g. ``007`` or ``+49…`` as text)."""
    if not isinstance(value, str):
        return value
    if value == "":
        return None
    if _NUMBER.fullmatch(value):
        number = float(value)
        return int(value) if re.fullmatch(r"-?\d+", value) else number
    return value


def table_to_records(header: list[str], rows: list[list[Any]], detect_numbers: bool
                     ) -> list[dict[str, Any]]:
    records = []
    for row in rows:
        values = (list(row) + [None] * len(header))[: len(header)]
        if detect_numbers:
            values = [coerce(v) for v in values]
        records.append({h: _plain(v) for h, v in zip(header, values, strict=True)})
    return records


def _plain(value: Any) -> Any:
    if isinstance(value, dt.datetime | dt.date | dt.time):
        return value.isoformat()
    return value


def extract_records(data: Any) -> list[dict[str, Any]]:
    if isinstance(data, list) and all(isinstance(r, dict) for r in data):
        return data
    if isinstance(data, dict):
        (only,) = data.values() if len(data) == 1 else (None,)
        if isinstance(only, list) and only and all(isinstance(r, dict) for r in only):
            return only  # e.g. {"rows": [...]} as produced for TOML
        return [data]  # a single object becomes a single row
    raise ConversionError(t("error.data_not_tabular"))


def flatten(record: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    flat: dict[str, Any] = {}
    for key, value in record.items():
        name = f"{prefix}{key}"
        if isinstance(value, dict):
            flat.update(flatten(value, name + "."))
        elif isinstance(value, list):
            flat[name] = json.dumps(value, ensure_ascii=False, default=str)
        else:
            flat[name] = value
    return flat


def records_to_table(data: Any) -> tuple[list[str], list[list[Any]]]:
    records = [flatten(r) for r in extract_records(data)]
    header: list[str] = []
    for r in records:
        header.extend(k for k in r if k not in header)
    return header, [[r.get(h) for h in header] for r in records]


# -- writing -------------------------------------------------------------------------------


def write_table(path: Path, fmt: str, header: list[str], rows: list[list[Any]],
                opts: dict[str, Any]) -> None:
    if fmt == "xlsx":
        from openpyxl import Workbook

        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Sheet1")
        ws.append(header)
        for row in rows:
            ws.append([_cell(v) for v in row])
        wb.save(path)
        return
    delimiter = "\t" if fmt == "tsv" else _DELIMITERS[opts.get("delimiter", "comma")]
    encoding = "utf-8-sig" if opts.get("excel_bom") else "utf-8"
    with path.open("w", encoding=encoding, newline="") as fh:
        writer = csv.writer(fh, delimiter=delimiter)
        writer.writerow(header)
        writer.writerows([[_csv_value(v) for v in row] for row in rows])


def _csv_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return value


def _cell(value: Any) -> Any:
    if value is None or isinstance(value, int | float | bool | dt.date | dt.datetime | dt.time):
        return value
    return str(value)


def write_structured(path: Path, fmt: str, data: Any, opts: dict[str, Any]) -> None:
    if fmt == "json":
        indent = int(opts.get("indent", 2)) or None
        text = json.dumps(data, indent=indent, ensure_ascii=False, default=str)
        path.write_text(text + "\n", encoding="utf-8")
    elif fmt == "yaml":
        import yaml

        text = yaml.safe_dump(_yaml_safe(data), allow_unicode=True, sort_keys=False)
        path.write_text(text, encoding="utf-8")
    else:
        import tomli_w

        if isinstance(data, list):
            data = {"rows": data}
        if not isinstance(data, dict):
            raise ConversionError(t("error.toml_structure"))
        path.write_text(tomli_w.dumps(_drop_none(data)), encoding="utf-8")


def _yaml_safe(data: Any) -> Any:
    if isinstance(data, dict):
        return {str(k): _yaml_safe(v) for k, v in data.items()}
    if isinstance(data, list):
        return [_yaml_safe(v) for v in data]
    return data


def _drop_none(data: Any) -> Any:
    """TOML has no null: drop such keys (inside lists an empty string keeps positions)."""
    if isinstance(data, dict):
        return {str(k): _drop_none(v) for k, v in data.items() if v is not None}
    if isinstance(data, list):
        return [_drop_none(v) if v is not None else "" for v in data]
    return data
